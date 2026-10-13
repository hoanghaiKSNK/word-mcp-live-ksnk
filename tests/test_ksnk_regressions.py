"""Synthetic regressions: no user document or running Word instance is touched."""
import asyncio
import json
import ntpath
from types import SimpleNamespace

import pytest
from docx import Document

from word_document_server.core import word_com, table_com
from word_document_server.tools import live_tools, live_read_tools
from word_document_server.utils import save_utils


class Docs:
    def __init__(self, docs): self.docs = docs
    @property
    def Count(self): return len(self.docs)
    def __call__(self, i): return self.docs[i - 1]
    def __iter__(self): return iter(self.docs)


def fake_app(paths):
    docs = [SimpleNamespace(Name=ntpath.basename(p), FullName=p) for p in paths]
    app = SimpleNamespace(Documents=Docs(docs), ActiveDocument=docs[0])
    for d in docs: d.Application = app
    return app


def test_full_path_never_matches_a_different_directory(monkeypatch):
    app = fake_app([r'D:\a\Draft.docx', r'D:\b\Draft.docx'])
    monkeypatch.setattr(word_com, '_iter_word_apps', lambda: [], raising=False)
    assert word_com.find_document(app, r'D:\b\Draft.docx') is app.Documents(2)


def test_missing_full_path_fails_even_when_basename_exists(monkeypatch):
    app = fake_app([r'D:\a\Draft.docx'])
    monkeypatch.setattr(word_com, '_iter_word_apps', lambda: [], raising=False)
    with pytest.raises(ValueError): word_com.find_document(app, r'D:\absent\Draft.docx')


def test_ambiguous_basename_fails_closed(monkeypatch):
    app = fake_app([r'D:\a\Draft.docx', r'D:\b\Draft.docx'])
    monkeypatch.setattr(word_com, '_iter_word_apps', lambda: [], raising=False)
    with pytest.raises(ValueError, match='[Aa]mbiguous'):
        word_com.find_document(app, 'Draft.docx')


def test_snapshot_keys_separate_same_named_documents():
    a = SimpleNamespace(Name='Draft.docx', FullName=r'D:\a\Draft.docx')
    b = SimpleNamespace(Name='Draft.docx', FullName=r'D:\b\Draft.docx')
    assert live_read_tools._doc_key(a) != live_read_tools._doc_key(b)


class Record:
    def __init__(self, recording=False): self.IsRecordingCustomRecord = recording; self.starts = 0; self.ends = 0
    def StartCustomRecord(self, name): self.starts += 1
    def EndCustomRecord(self): self.ends += 1


def test_existing_user_undo_record_is_not_closed():
    rec = Record(True); app = SimpleNamespace(UndoRecord=rec, ScreenUpdating=True)
    with pytest.raises(RuntimeError):
        with word_com.undo_record(app, 'test'): pass
    assert rec.starts == rec.ends == 0


def test_screen_updating_is_restored_on_failure():
    app = SimpleNamespace(UndoRecord=Record(), ScreenUpdating=True)
    with pytest.raises(ValueError):
        with word_com.undo_record(app, 'test'):
            assert app.ScreenUpdating is False
            raise ValueError('synthetic failure')
    assert app.ScreenUpdating is True
    assert app.UndoRecord.ends == 1


def test_failed_atomic_save_keeps_old_destination(tmp_path, monkeypatch):
    import docx.document
    old_save = docx.document.Document.save
    monkeypatch.setattr(docx.document.Document, 'save', old_save)
    target = tmp_path / 'saved.docx'; old = Document(); old.add_paragraph('ORIGINAL'); old.save(target)
    before = target.read_bytes(); new = Document(); new.add_paragraph('NEW')
    save_utils.install_save_hook()
    def fail_replace(*args): raise OSError('synthetic replace failure')
    monkeypatch.setattr('os.replace', fail_replace)
    with pytest.raises(OSError): new.save(target)
    assert target.read_bytes() == before


def test_save_as_preserves_source_comments_not_destination(tmp_path, monkeypatch):
    import docx.document
    old_save = docx.document.Document.save
    monkeypatch.setattr(docx.document.Document, 'save', old_save)
    source = Document(); run = source.add_paragraph('KSNK source').runs[0]
    source.add_comment(run, text='SOURCE COMMENT', author='Synthetic Author')
    source_path = tmp_path / 'source.docx'; source.save(source_path)
    dest = Document(); run2 = dest.add_paragraph('Other document').runs[0]
    dest.add_comment(run2, text='DESTINATION COMMENT', author='Other Author')
    target = tmp_path / 'target.docx'; dest.save(target)
    save_utils.install_save_hook(); Document(source_path).save(target)
    loaded = Document(target)
    assert [c.text for c in loaded.comments] == ['SOURCE COMMENT']


def test_save_plain_source_does_not_inherit_destination_comments(tmp_path, monkeypatch):
    import docx.document
    monkeypatch.setattr(docx.document.Document, 'save', docx.document.Document.save)
    source=Document(); source.add_paragraph('Plain source')
    other=Document(); run=other.add_paragraph('Other source').runs[0]
    other.add_comment(run,text='UNRELATED COMMENT',author='Other')
    target=tmp_path/'destination.docx'; other.save(target)
    save_utils.install_save_hook(); source.save(target)
    assert len(Document(target).comments)==0
    assert len(list((tmp_path/'.word-mcp-backups').glob('destination.docx.*.bak')))==1


def test_edit_session_restores_author_and_tracking_on_setup_failure():
    class App:
        UndoRecord=Record(); ScreenUpdating=True
        def __init__(self): self.author='Original'; self.fail=True
        @property
        def UserName(self): return self.author
        @UserName.setter
        def UserName(self,value):
            if self.fail:
                self.fail=False
                raise RuntimeError('synthetic setup failure')
            self.author=value
    app=App(); doc=SimpleNamespace(TrackRevisions=False)
    with pytest.raises(RuntimeError):
        with word_com.edit_session(app,doc,'test',True): pytest.fail('must not edit')
    assert doc.TrackRevisions is False and app.UserName=='Original' and app.ScreenUpdating


class Cell:
    def __init__(self, row, col):
        self.Range = SimpleNamespace(Text='old\r\x07', Start=(row*10+col), End=(row*10+col+4), Revisions=SimpleNamespace(Count=0))


class Table:
    def __init__(self, rows=1, cols=2):
        self.Rows = SimpleNamespace(Count=rows); self.Columns = SimpleNamespace(Count=cols)
        self.cells = {(r,c): Cell(r,c) for r in range(1, rows+1) for c in range(1,cols+1)}
    def Cell(self,r,c): return self.cells[r,c]


def test_bulk_write_rejects_invalid_dimensions_before_any_write():
    t=Table(); before=[c.Range.Text for c in t.cells.values()]
    with pytest.raises(ValueError): table_com.set_range(t, [['one','two','OVERFLOW']])
    assert [c.Range.Text for c in t.cells.values()] == before


@pytest.mark.parametrize('value', ['bad\x07cell','bad\x00text'])
def test_bulk_write_rejects_control_bytes_before_any_write(value):
    t=Table(); before=[c.Range.Text for c in t.cells.values()]
    with pytest.raises(ValueError): table_com.set_row(t, 1, ['one',value])
    assert [c.Range.Text for c in t.cells.values()] == before


def test_preflight_merged_cell_aliases():
    t=Table(); t.cells[1,2]=t.cells[1,1]
    with pytest.raises(ValueError, match='[Mm]erged|[Aa]lias'):
        table_com.set_row(t,1,['one','two'])
    assert t.cells[1,1].Range.Text == 'old\r\x07'
